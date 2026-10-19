#!/usr/bin/env python3
"""文字起こしのテスト用の .pptx と .xlsx を作る(作り直すときだけ使う。python-pptx と openpyxl が要る)。
使い方: python3 tests/transcribe/make-fixtures.py
Excel が保存したファイルに近づけるため、式の計算結果(<v>)とふりがな(<rPh>)を後から書き足す。
"""
import io, os, re, struct, zipfile, zlib
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
import openpyxl, datetime

D = os.path.join(os.path.dirname(__file__), 'fixtures')

def png():
    raw = b'\x00\xff\x00\x00'
    def chunk(t, d): return struct.pack('>I', len(d)) + t + d + struct.pack('>I', zlib.crc32(t + d) & 0xffffffff)
    return b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('>IIBBBBB', 1, 1, 8, 2, 0, 0, 0)) + chunk(b'IDAT', zlib.compress(raw)) + chunk(b'IEND', b'')

# ---- PowerPoint
prs = Presentation()
s1 = prs.slides.add_slide(prs.slide_layouts[1])          # タイトルと本文
s1.shapes.title.text = '送料無料バッジの要件'
tf = s1.placeholders[1].text_frame
tf.text = '対象: 商品詳細ページ'
p = tf.add_paragraph(); p.text = '5,000 円以上で表示'; p.level = 1
p = tf.add_paragraph(); p.text = '会員ランクは見ない'; p.level = 1
p = tf.add_paragraph(); p.text = 'しきい値は Commerce 側で持つ | OCC で返す'
s1.notes_slide.notes_text_frame.text = 'しきい値の根拠は営業部に確認済み\n変更は管理画面から'

s2 = prs.slides.add_slide(prs.slide_layouts[5])          # タイトルのみ
s2.shapes.title.text = 'OCC の項目'
rows, cols = 3, 3
tbl = s2.shapes.add_table(rows, cols, Inches(0.5), Inches(1.5), Inches(8), Inches(1.5)).table
for r, vals in enumerate([['項目', '型', '説明'], ['freeShippingThreshold', 'Price', 'しきい値\n(税込)'], ['eligible', 'boolean', '表示するか']]):
    for c, v in enumerate(vals):
        tbl.cell(r, c).text = v
grp = s2.shapes.add_group_shape()
box = grp.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(4), Inches(2), Inches(1)); box.text = 'グループ内の注記'
s2.shapes.add_picture(io.BytesIO(png()), Inches(6), Inches(4), Inches(1), Inches(1))
s2.shapes[-1]._element.nvPicPr.cNvPr.set('descr', '画面イメージ(バッジの位置)')

s3 = prs.slides.add_slide(prs.slide_layouts[0])          # タイトルスライド
s3.shapes.title.text = '引き継ぎ資料'
s3.placeholders[1].text = '2026 年 10 月 フロントエンド'

s4 = prs.slides.add_slide(prs.slide_layouts[5])
s4.shapes.title.text = '旧案(使わない)'
s4._element.set('show', '0')
cd = CategoryChartData(); cd.categories = ['A', 'B']; cd.add_series('件数', (1, 2))
s4.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(4), Inches(3), cd)

s5 = prs.slides.add_slide(prs.slide_layouts[6])          # 白紙

# 表示順を slide3 → slide1 → slide2 → slide4 → slide5 にする(ファイル名の順と違う)
lst = prs.slides._sldIdLst; ids = list(lst)
for el in ids: lst.remove(el)
for el in [ids[2], ids[0], ids[1], ids[3], ids[4]]: lst.append(el)
prs.save(os.path.join(D, 'sample.pptx'))

# SmartArt(python-pptx では作れないので、図形とデータの部品を書き足す)。非表示のスライド(slide4.xml)に置く
zin = zipfile.ZipFile(os.path.join(D, 'sample.pptx')); items = {n: zin.read(n) for n in zin.namelist()}; zin.close()
A_ = 'http://schemas.openxmlformats.org/drawingml/2006/main'
frame = ('<p:graphicFrame><p:nvGraphicFramePr><p:cNvPr id="90" name="SmartArt"/><p:cNvGraphicFramePr/><p:nvPr/></p:nvGraphicFramePr>'
         '<p:xfrm><a:off x="0" y="0"/><a:ext cx="100" cy="100"/></p:xfrm><a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/diagram">'
         '<dgm:relIds xmlns:dgm="http://schemas.openxmlformats.org/drawingml/2006/diagram" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:dm="rIdDgm1" r:lo="rIdDgm2" r:qs="rIdDgm3" r:cs="rIdDgm4"/>'
         '</a:graphicData></a:graphic></p:graphicFrame>')
sl = items['ppt/slides/slide4.xml'].decode()
sl = sl.replace('</p:spTree>', frame + '</p:spTree>', 1); items['ppt/slides/slide4.xml'] = sl.encode()
def pt(i, text, typ=None):
    t = ' type="%s"' % typ if typ else ''
    body = '<dgm:t><a:bodyPr/><a:p><a:r><a:t>%s</a:t></a:r></a:p></dgm:t>' % text if text else ''
    return '<dgm:pt modelId="%d"%s>%s</dgm:pt>' % (i, t, body)
data = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><dgm:dataModel xmlns:dgm="http://schemas.openxmlformats.org/drawingml/2006/diagram" xmlns:a="%s"><dgm:ptLst>' % A_
        + pt(0, None, 'doc') + pt(1, '要件') + pt(2, '設計') + pt(3, '並べ替えの印', 'parTrans') + pt(4, 'リリース') + '</dgm:ptLst></dgm:dataModel>')
items['ppt/diagrams/data1.xml'] = data.encode('utf-8')
rels = items['ppt/slides/_rels/slide4.xml.rels'].decode()
rels = rels.replace('</Relationships>', '<Relationship Id="rIdDgm1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/diagramData" Target="../diagrams/data1.xml"/></Relationships>')
items['ppt/slides/_rels/slide4.xml.rels'] = rels.encode()
ct = items['[Content_Types].xml'].decode()
ct = ct.replace('</Types>', '<Override PartName="/ppt/diagrams/data1.xml" ContentType="application/vnd.openxmlformats-officedocument.drawingml.diagramData+xml"/></Types>')
items['[Content_Types].xml'] = ct.encode()
with zipfile.ZipFile(os.path.join(D, 'sample.pptx'), 'w', zipfile.ZIP_DEFLATED) as zout:
    for n, b in items.items(): zout.writestr(n, b)

# ---- Excel
wb = openpyxl.Workbook()
ws = wb.active; ws.title = 'API一覧'
ws['A1'] = 'OCC API 一覧'; ws.merge_cells('A1:D1')
for c, v in zip('ABCDF', ['No', 'API', 'メソッド', '件数', '備考']): ws[c + '3'] = v
ws.append([])
data = [(1, '/users/{userId}/carts', 'POST', 120, '改行\nあり'), (2, '/products/{code}', 'GET', 3400, 'a|b'), (3, '/cms/pages', 'GET', 0.1, None)]
for i, (n, api, m, cnt, note) in enumerate(data, 4):
    ws['A%d' % i] = n; ws['B%d' % i] = api; ws['C%d' % i] = m; ws['D%d' % i] = cnt
    if note: ws['F%d' % i] = note
ws['C8'] = '合計'; ws['D8'] = '=SUM(D4:D6)'
ws['A10'] = '更新日'; ws['B10'] = datetime.date(2026, 10, 19); ws['B10'].number_format = 'yyyy/m/d'
ws['A11'] = '時刻'; ws['B11'] = datetime.datetime(2026, 10, 19, 9, 30); ws['B11'].number_format = 'yyyy/m/d h:mm'
ws['A12'] = '公開'; ws['B12'] = True
ws['A13'] = '税込'; ws['B13'] = 1500; ws['B13'].number_format = '#,##0"円"'
ws2 = wb.create_sheet('メモ'); ws2['A1'] = '非表示のメモ'; ws2.sheet_state = 'hidden'
wb.create_sheet('空')
path = os.path.join(D, 'sample.xlsx'); wb.save(path)

# 計算結果とふりがなを足す(Excel で保存したときと同じ形にする)
zin = zipfile.ZipFile(path); items = {n: zin.read(n) for n in zin.namelist()}; zin.close()
sheet = items['xl/worksheets/sheet1.xml'].decode()
sheet = sheet.replace('<f>SUM(D4:D6)</f><v></v>', '<f>SUM(D4:D6)</f><v>3520.1</v>').replace('<f>SUM(D4:D6)</f></c>', '<f>SUM(D4:D6)</f><v>3520.1</v></c>')
assert '3520.1' in sheet
items['xl/worksheets/sheet1.xml'] = sheet.encode()
# Excel は文字を共有文字列(sharedStrings.xml)に置く。F5 だけはセル内の文字(inlineStr)のまま残す
from html import unescape
from xml.sax.saxutils import escape
strings = []
def to_shared(m):
    ref, text = m.group(1), unescape(m.group(2))
    if ref == 'F5':
        return m.group(0)
    strings.append(text)
    return '<c r="%s" t="s"><v>%d</v></c>' % (ref, len(strings) - 1)
sheet = re.sub(r'<c r="([A-Z]+[0-9]+)" t="inlineStr"><is><t>(.*?)</t></is></c>', to_shared, sheet, flags=re.S)
items['xl/worksheets/sheet1.xml'] = sheet.encode()
def si(text):
    if text == '更新日':   # ふりがな(rPh)は文字起こしに出さない
        return '<si><t>更新日</t><rPh sb="0" eb="3"><t>コウシンビ</t></rPh><phoneticPr fontId="1"/></si>'
    if text == 'メソッド':  # 書式の違う部分(r)に分かれた文字
        return '<si><r><t>メソ</t></r><r><rPr><b/></rPr><t>ッド</t></r></si>'
    return '<si><t xml:space="preserve">%s</t></si>' % escape(text)
items['xl/sharedStrings.xml'] = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="%d" uniqueCount="%d">' % (len(strings), len(strings)) + ''.join(si(t) for t in strings) + '</sst>').encode('utf-8')
rels = items['xl/_rels/workbook.xml.rels'].decode()
rels = rels.replace('</Relationships>', '<Relationship Id="rIdSst" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings" Target="sharedStrings.xml"/></Relationships>')
items['xl/_rels/workbook.xml.rels'] = rels.encode()
ct = items['[Content_Types].xml'].decode()
ct = ct.replace('</Types>', '<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/></Types>')
items['[Content_Types].xml'] = ct.encode()
with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as zout:
    for n, b in items.items(): zout.writestr(n, b)
print('作った:', os.listdir(D))
